"""Проверка структуры файла"""
import pandas as pd
from openpyxl import load_workbook

input_file = "output/your_file_modified.xlsx"

# Проверка через pandas
print("=== Проверка через pandas ===")
df = pd.read_excel(input_file, header=None)
print(f"Размер DataFrame: {df.shape}")
print(f"\nПервые 15 строк:")
print(df.head(15))
print(f"\nКолонки: {df.columns.tolist()}")

# Проверка через openpyxl
print("\n=== Проверка через openpyxl ===")
wb = load_workbook(input_file)
ws = wb.active
print(f"Максимальная строка: {ws.max_row}")
print(f"Максимальная колонка: {ws.max_column}")

print("\nПервые 15 строк (первые 10 колонок):")
for row_num in range(1, 16):
    row_data = []
    for col_num in range(1, min(11, ws.max_column + 1)):
        cell = ws.cell(row_num, col_num)
        row_data.append(str(cell.value)[:30] if cell.value else '')
    print(f"Строка {row_num:2d}: {row_data}")
